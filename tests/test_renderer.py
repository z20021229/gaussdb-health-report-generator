from pathlib import Path

from docx import Document
from PIL import Image

from src.analyzer import analyze_inspection_data
from src.archive_extractor import extract_package
from src.evidence_image_builder import build_evidence_images
from src.inspection_parser import parse_inspection_files
from src.renderer import render_docx


BANNED_WORDS = [
    "workdir",
    "output",
    "raw_sections",
    "evidence_images",
    "extracted_manifest",
    "evidence_images_manifest",
    "inspection_data.generated.yaml",
    ".yaml",
    "raw_detail_path",
    "source_file",
    "manifest",
    "parser",
    "analyzer",
    "renderer",
    "程序解析",
    "代码发现",
    "YAML",
]

EXPECTED_CHAPTERS = [
    "第一章 总结",
    "第二章 系统概况",
    "2.1. 操作系统版本检查",
    "2.2. 数据库版本检查",
    "2.3. cpu 核数信息",
    "第三章 总体情况",
    "3.1. 集群运行情况",
    "3.2. 磁盘空间概况",
    "第四章 高可用检查",
    "4.1. 集群高可用状态检查",
    "4.2. CPU 一天使用信息",
    "4.3. 数据库运行状态",
    "4.4. 复制槽状态",
    "第五章 参数检查",
    "5.1. 函数运行状态检查",
    "5.2. 数据库信息检查",
    "5.3. gs_collector 信息收集",
    "第六章 系统管理维护",
    "6.1. 大表检查",
    "6.2. 未使用的索引",
    "6.3. 索引建议",
    "6.4. 表膨胀检查",
]


def test_render_docx_uses_reference_chapter_structure_without_appendix(tmp_path: Path) -> None:
    image_path = _create_image(tmp_path / "evidence" / "os.png")
    data = _base_report_data(
        evidence_items=[
            {
                "type": "section_text",
                "title": "操作系统信息 - 第 1 页",
                "section_name": "操作系统信息",
                "source_file": "sample/inspection_rec.txt",
                "image_path": str(image_path),
                "status": "success",
                "error": "",
            }
        ]
    )
    output_path = tmp_path / "report.docx"

    render_docx(data, output_path)

    document = Document(output_path)
    text = _document_text(document)
    for chapter in EXPECTED_CHAPTERS:
        assert chapter in text
    assert "附录" not in text
    assert "风险级问题明细" not in text
    assert "关注级问题明细" not in text
    assert "整改建议汇总" not in text
    for banned in BANNED_WORDS:
        assert banned not in text
    assert len(document.inline_shapes) == 1


def test_render_docx_marks_disk_25_percent_as_good(tmp_path: Path) -> None:
    data = _base_report_data()
    data["nodes"][0]["disks"][0]["use_percent"] = 25
    data["nodes"][0]["disk_summary"]["max_disk_use_percent"] = 25
    data["cluster"]["resource_summary"]["cluster_max_disk_use_percent"] = 25
    output_path = tmp_path / "disk.docx"

    render_docx(data, output_path)

    text = _document_text(Document(output_path))
    disk_section = text.split("3.2. 磁盘空间概况", 1)[1]
    assert "整体状态为良好" in disk_section
    assert "磁盘空间充足" in disk_section
    assert "风险" not in disk_section[:180]


def test_database_version_without_screenshot_does_not_borrow_database_info_image(tmp_path: Path) -> None:
    db_info_image = _create_image(tmp_path / "evidence" / "database_info.png")
    data = _base_report_data(
        evidence_items=[
            {
                "type": "section_text",
                "title": "数据库信息检查 - 第 1 页",
                "section_name": "数据库信息检查",
                "source_file": "sample/inspection_rec.txt",
                "image_path": str(db_info_image),
                "status": "success",
                "error": "",
            }
        ]
    )
    output_path = tmp_path / "database-version.docx"

    render_docx(data, output_path)

    document = Document(output_path)
    text = _document_text(document)
    db_version_section = text.rsplit("2.2. 数据库版本检查", 1)[1].split("2.3. cpu 核数信息", 1)[0]
    assert "检查结果未采集。" in db_version_section
    assert len(document.inline_shapes) == 1


def test_empty_fatal_and_panic_logs_do_not_render_as_exceptions(tmp_path: Path) -> None:
    data = _base_report_data()
    output_path = tmp_path / "logs.docx"

    render_docx(data, output_path)

    text = _document_text(Document(output_path))
    assert "没有严重的异常和告警。" in text
    assert "日志文件缺失" not in text
    assert "发现需关注日志信息" not in text


def test_gs_check_and_table_bloat_stay_in_their_own_chapters(tmp_path: Path) -> None:
    data = _base_report_data()
    output_path = tmp_path / "sections.docx"

    render_docx(data, output_path)

    text = _document_text(Document(output_path))
    chapter_one = text.rsplit("第一章 总结", 1)[1].split("第二章 系统概况", 1)[0]
    chapter_five = text.rsplit("第五章 参数检查", 1)[1].split("第六章 系统管理维护", 1)[0]
    chapter_six = text.rsplit("第六章 系统管理维护", 1)[1]
    assert "CheckDirPermissions" not in chapter_one
    assert "CheckDirPermissions" in chapter_five
    assert "表膨胀检查发现" in chapter_six


def test_render_real_sample_docx_without_engineering_terms(tmp_path: Path) -> None:
    root = Path(__file__).resolve().parents[1]
    sample_input = root / "samples" / "收益所有人.rar"
    workdir = tmp_path / "workdir"
    output_dir = tmp_path / "report_out"
    output_dir.mkdir(parents=True, exist_ok=True)

    manifest = extract_package(str(sample_input), str(workdir))
    inspection_files = manifest.get("files", {}).get("inspection_rec", [])
    data = parse_inspection_files(inspection_files, manifest=manifest, output_dir=output_dir)
    data = analyze_inspection_data(data)
    evidence = build_evidence_images(data, manifest, str(output_dir))
    data["evidence_images"] = {
        "manifest_path": evidence["manifest_path"],
        "evidence_images_dir": evidence["evidence_images_dir"],
        "items": evidence["items"],
    }

    output_path = tmp_path / "real.docx"
    render_docx(data, output_path)

    document = Document(output_path)
    text = _document_text(document)
    for chapter in EXPECTED_CHAPTERS:
        assert chapter in text
    assert "附录" not in text
    for banned in BANNED_WORDS:
        assert banned not in text
    assert len(document.inline_shapes) > 0


def _base_report_data(*, evidence_items: list[dict] | None = None) -> dict:
    return {
        "report": {
            "title": "GaussDB 数据库健康诊断报告",
            "inspector": "未提供",
            "inspection_date": "未采集",
        },
        "nodes": [
            {
                "ip": "10.0.0.1",
                "hostname": "node-a",
                "role": "未采集",
                "os_version": "openEuler 22.03",
                "architecture": "x86_64",
                "cpu_model": "CPU Demo",
                "cpu_cores": 16,
                "memory": {"usage_percent": 50.0},
                "disks": [
                    {
                        "filesystem": "/dev/vda1",
                        "size": "500G",
                        "used": "125G",
                        "avail": "375G",
                        "use_percent": 25,
                        "mounted_on": "/data",
                    }
                ],
                "disk_summary": {"max_disk_use_percent": 25},
                "cpu_daily": {
                    "avg_user": 8.2,
                    "avg_system": 2.3,
                    "avg_iowait": 0.4,
                    "avg_idle": 89.1,
                    "min_idle": 80.0,
                    "max_iowait": 1.2,
                },
            }
        ],
        "cluster": {
            "cluster_status": "未采集",
            "overall_status": "关注",
            "ha_status": [{"client_addr": "10.0.0.2", "sync_state": "Sync", "pg_xlog_location_diff": "0"}],
            "replication_slots": [{"slot_name": "slot_demo", "slot_type": "physical", "active": "true", "delay_lsn": "0"}],
            "gs_check_summary": {
                "ok_count": 3,
                "ng_count": 1,
                "na_count": 0,
                "unknown_count": 0,
                "ng_display_items": [
                    {
                        "check_name": "CheckDirPermissions",
                        "status": "NG",
                        "detail_summary": "目录权限需要进一步核查",
                        "suggestion": "建议根据安全规范核查目录权限，避免权限过宽。",
                    }
                ],
            },
            "resource_summary": {
                "node_count": 1,
                "cluster_max_disk_use_percent": 25,
                "cluster_min_cpu_idle": 80.0,
                "cluster_max_iowait": 1.2,
            },
        },
        "database": {
            "version": "GaussDB 5.x",
            "running_status": {"records": [{"is_in_recovery": "false"}]},
            "databases": [{"datname": "demo", "readable_size": "10GB"}],
            "large_tables": [{"relname": "large_table"}],
            "index_suggestions": [{"tablename": "demo_table"}],
            "unused_indexes": [],
            "table_bloat": [{"schemaname": "public", "relname": "t_demo", "dead_rate": "25%"}],
        },
        "logs": {
            "fatal_log_summary": "未发现 fatal 异常日志",
            "panic_log_summary": "未发现 panic 异常日志",
        },
        "evidence_images": {"items": evidence_items or []},
    }


def _create_image(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (200, 100), color="white").save(path)
    return path


def _document_text(document: Document) -> str:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.extend(cell.text for cell in row.cells)
    return "\n".join(paragraphs)
