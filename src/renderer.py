"""Render a formal customer-facing GaussDB health diagnosis Word report."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


IMAGE_WIDTH = Inches(6.0)
NOT_COLLECTED = "未采集"


CHAPTER_TOC = [
    ("第一章 总结", "1"),
    ("第二章 系统概况", "2"),
    ("2.1. 操作系统版本检查", "2"),
    ("2.2. 数据库版本检查", "3"),
    ("2.3. cpu 核数信息", "3"),
    ("第三章 总体情况", "4"),
    ("3.1. 集群运行情况", "4"),
    ("3.2. 磁盘空间概况", "4"),
    ("第四章 高可用检查", "5"),
    ("4.1. 集群高可用状态检查", "5"),
    ("4.2. CPU 一天使用信息", "6"),
    ("4.3. 数据库运行状态", "6"),
    ("4.4. 复制槽状态", "7"),
    ("第五章 参数检查", "8"),
    ("5.1. 函数运行状态检查", "8"),
    ("5.2. 数据库信息检查", "8"),
    ("5.3. gs_collector 信息收集", "9"),
    ("第六章 系统管理维护", "10"),
    ("6.1. 大表检查", "10"),
    ("6.2. 未使用的索引", "11"),
    ("6.3. 索引建议", "11"),
    ("6.4. 表膨胀检查", "12"),
]


def render_docx(data: dict[str, Any], output_path: Path) -> Path:
    """Create the final report docx and return the saved path."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _configure_document(document)

    _render_cover(document, data)
    _render_inspection_date_page(document, data)
    _render_toc(document)
    _render_chapter_one(document, data)
    _render_chapter_two(document, data)
    _render_chapter_three(document, data)
    _render_chapter_four(document, data)
    _render_chapter_five(document, data)
    _render_chapter_six(document, data)

    try:
        document.save(output_path)
        return output_path
    except PermissionError:
        timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
        for index in range(1, 20):
            suffix = f"_latest_{timestamp}" if index == 1 else f"_latest_{timestamp}_{index}"
            fallback_path = output_path.with_name(f"{output_path.stem}{suffix}{output_path.suffix}")
            try:
                document.save(fallback_path)
                return fallback_path
            except PermissionError:
                continue
        raise


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)


def _render_cover(document: Document, data: dict[str, Any]) -> None:
    report = data.get("report", {})
    title = _text(report.get("title"), "GaussDB 数据库健康诊断报告")
    inspector = _text(report.get("inspector"), "未提供")

    document.add_paragraph("")
    document.add_paragraph("")
    title_para = document.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(24)

    for _ in range(8):
        document.add_paragraph("")

    inspector_para = document.add_paragraph()
    inspector_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inspector_para.add_run(f"巡检人：{inspector}")
    document.add_page_break()


def _render_inspection_date_page(document: Document, data: dict[str, Any]) -> None:
    _add_heading(document, "巡检日期", level=1)
    document.add_paragraph(f"巡检日期：{_inspection_date(data)}")
    document.add_page_break()


def _render_toc(document: Document) -> None:
    _add_heading(document, "目录", level=1)
    for title, page in CHAPTER_TOC:
        dots = "." * max(4, 42 - len(title))
        document.add_paragraph(f"{title}{dots}{page}")
    document.add_page_break()


def _render_chapter_one(document: Document, data: dict[str, Any]) -> None:
    _add_heading(document, "第一章 总结", level=1)
    document.add_paragraph("本次巡检的总结如下：")
    for index, line in enumerate(_summary_lines(data), start=1):
        document.add_paragraph(f"{index}. {line}")
    document.add_page_break()


def _render_chapter_two(document: Document, data: dict[str, Any]) -> None:
    _add_heading(document, "第二章 系统概况", level=1)

    _render_check_section(
        document,
        data,
        "2.1. 操作系统版本检查",
        ["操作系统信息"],
        _os_conclusion(data),
        exact_only=True,
    )

    database = data.get("database", {})
    _add_heading(document, "2.2. 数据库版本检查", level=2)
    document.add_paragraph(f"数据库版本：{_text(database.get('version'))}")
    _render_section_evidence(
        document,
        data,
        ["数据库版本检查"],
        exact_only=True,
        empty_message="检查结果未采集。",
    )
    _add_conclusion(document, _database_version_conclusion(data))

    _render_check_section(
        document,
        data,
        "2.3. cpu 核数信息",
        ["CPU核数信息", "CPU型号信息"],
        _cpu_conclusion(data),
        exact_only=True,
    )
    document.add_page_break()


def _render_chapter_three(document: Document, data: dict[str, Any]) -> None:
    _add_heading(document, "第三章 总体情况", level=1)

    _add_heading(document, "3.1. 集群运行情况", level=2)
    cluster_status = _text(data.get("cluster", {}).get("cluster_status"))
    if cluster_status == NOT_COLLECTED:
        document.add_paragraph("本次巡检未获取到集群整体状态检查结果，建议后续补充 gs_om -t status 等集群状态采集。")
    else:
        document.add_paragraph(f"集群整体状态：{cluster_status}")
    _render_section_evidence(
        document,
        data,
        ["集群运行情况", "集群状态", "集群运行"],
        exact_only=True,
        empty_message="检查结果未采集。",
    )
    _add_conclusion(document, _cluster_status_conclusion(data))

    _render_check_section(
        document,
        data,
        "3.2. 磁盘空间概况",
        ["磁盘空间大小", "DN占用空间大小", "ETCD占用空间大小"],
        _disk_conclusion(data),
        exact_only=True,
    )
    document.add_page_break()


def _render_chapter_four(document: Document, data: dict[str, Any]) -> None:
    _add_heading(document, "第四章 高可用检查", level=1)
    _render_check_section(
        document,
        data,
        "4.1. 集群高可用状态检查",
        ["集群高可用状态检查"],
        _ha_conclusion(data),
        exact_only=True,
    )
    _render_check_section(
        document,
        data,
        "4.2. CPU 一天使用信息",
        ["CPU近一天使用情况"],
        _cpu_daily_conclusion(data),
        exact_only=True,
    )
    _render_check_section(
        document,
        data,
        "4.3. 数据库运行状态",
        ["数据库运行状态检查"],
        _running_status_conclusion(data),
        exact_only=True,
    )
    _render_check_section(
        document,
        data,
        "4.4. 复制槽状态",
        ["复制槽状态检查"],
        _replication_slot_conclusion(data),
        exact_only=True,
    )
    document.add_page_break()


def _render_chapter_five(document: Document, data: dict[str, Any]) -> None:
    _add_heading(document, "第五章 参数检查", level=1)
    _render_check_section(
        document,
        data,
        "5.1. 函数运行状态检查",
        ["函数运行状态检查"],
        "函数运行状态检查完成，未发现明显异常。",
        exact_only=True,
    )
    _render_check_section(
        document,
        data,
        "5.2. 数据库信息检查",
        ["数据库信息检查"],
        _database_info_conclusion(data),
        exact_only=True,
    )

    _add_heading(document, "5.3. gs_collector 信息收集", level=2)
    _render_section_evidence(document, data, ["gs_collector信息", "gs_collector 信息"], exact_only=True)
    _render_gs_check_summary(document, data)
    _add_conclusion(document, _gs_check_conclusion(data))
    document.add_page_break()


def _render_chapter_six(document: Document, data: dict[str, Any]) -> None:
    _add_heading(document, "第六章 系统管理维护", level=1)
    _render_check_section(
        document,
        data,
        "6.1. 大表检查",
        ["大表检查"],
        _large_table_conclusion(data),
        exact_only=True,
    )
    _render_check_section(
        document,
        data,
        "6.2. 未使用的索引",
        ["未使用的索引"],
        _unused_index_conclusion(data),
        exact_only=True,
    )
    _render_check_section(
        document,
        data,
        "6.3. 索引建议",
        ["索引建议"],
        _index_suggestion_conclusion(data),
        exact_only=True,
    )
    _render_check_section(
        document,
        data,
        "6.4. 表膨胀检查",
        ["表膨胀检查"],
        _table_bloat_conclusion(data),
        exact_only=True,
    )


def _render_check_section(
    document: Document,
    data: dict[str, Any],
    heading: str,
    section_names: list[str],
    conclusion: str,
    *,
    exact_only: bool = True,
) -> None:
    _add_heading(document, heading, level=2)
    _render_section_evidence(document, data, section_names, exact_only=exact_only)
    _add_conclusion(document, conclusion)


def _render_section_evidence(
    document: Document,
    data: dict[str, Any],
    section_names: list[str],
    *,
    exact_only: bool = True,
    empty_message: str = "检查结果未采集。",
) -> int:
    items = _matching_section_images(data, section_names, exact_only=exact_only)
    if not items:
        document.add_paragraph(empty_message)
        return 0

    for item in items:
        document.add_paragraph("检查结果：")
        image_path = Path(str(item.get("image_path", "")))
        try:
            document.add_picture(str(image_path), width=IMAGE_WIDTH)
        except Exception:
            document.add_paragraph("检查结果图片插入失败。")
    return len(items)


def _matching_section_images(
    data: dict[str, Any],
    section_names: list[str],
    *,
    exact_only: bool = True,
) -> list[dict[str, Any]]:
    evidence = data.get("evidence_images", {})
    items = evidence.get("items", [])
    matched: list[dict[str, Any]] = []
    targets = [name.lower().replace(" ", "") for name in section_names]
    for item in items:
        if item.get("type") != "section_text" or item.get("status") != "success":
            continue
        image_path = Path(str(item.get("image_path", "")))
        if not image_path.exists():
            continue
        section_name = str(item.get("section_name") or "")
        normalized = section_name.lower().replace(" ", "")
        if exact_only:
            ok = normalized in targets
        else:
            ok = any(target in normalized or normalized in target for target in targets)
        if ok:
            matched.append(item)
    return matched


def _render_gs_check_summary(document: Document, data: dict[str, Any]) -> None:
    summary = data.get("cluster", {}).get("gs_check_summary", {})
    paragraph = document.add_paragraph()
    paragraph.add_run("gs_check 巡检结果：").bold = True
    paragraph.add_run(
        f"OK {summary.get('ok_count', 0)} 项，"
        f"NG {summary.get('ng_count', 0)} 项，"
        f"NA {summary.get('na_count', 0)} 项，"
        f"UNKNOWN {summary.get('unknown_count', 0)} 项。"
    )

    ng_items = summary.get("ng_display_items") or summary.get("ng_items") or []
    if not ng_items:
        return

    for index, item in enumerate(ng_items, start=1):
        check_name = _text(item.get("check_name"))
        detail = _short_text(item.get("detail_summary"), 120)
        suggestion = _text(item.get("suggestion"), _gs_check_suggestion(check_name))
        document.add_paragraph(f"{index}. {check_name}：状态 NG；{detail}；{suggestion}")


def _add_heading(document: Document, text: str, *, level: int) -> None:
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        run.bold = True
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(16 if level == 1 else 13)


def _add_conclusion(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run("结论：").bold = True
    paragraph.add_run(_text(text))


def _summary_lines(data: dict[str, Any]) -> list[str]:
    cluster = data.get("cluster", {})
    database = data.get("database", {})
    return [
        f"数据库整体状态{_overall_status_phrase(cluster.get('overall_status'))}。",
        f"日志：{_log_sentence(data)}",
        f"实例状态：{_running_status_conclusion(data)}",
        f"磁盘：{_disk_conclusion(data)}",
        f"standby：{_ha_conclusion(data)}",
        f"内存：{_memory_conclusion(data)}",
        f"数据库信息：{_database_info_conclusion(data)}",
        f"巡检发现：{_major_findings_sentence(data, database)}",
    ]


def _overall_status_phrase(status: Any) -> str:
    status_text = _text(status, "良好")
    if status_text == "良好":
        return "良好"
    if status_text == "风险":
        return "存在风险项"
    if status_text == "关注":
        return "存在需关注项"
    return f"为{status_text}"


def _log_sentence(data: dict[str, Any]) -> str:
    logs = data.get("logs", {})
    fatal = _text(logs.get("fatal_log_summary"))
    panic = _text(logs.get("panic_log_summary"))
    combined = f"{fatal}；{panic}"
    if "未发现" in combined and "异常日志" in combined:
        return "没有严重的异常和告警。"
    if "未发现日志文件" in combined:
        return "未获取到对应日志文件。"
    if fatal == NOT_COLLECTED and panic == NOT_COLLECTED:
        return "未获取到对应日志文件。"
    return f"发现需关注日志信息，建议结合数据库运行时段核查。{_short_text(combined, 120)}"


def _running_status_conclusion(data: dict[str, Any]) -> str:
    records = data.get("database", {}).get("running_status", {}).get("records", [])
    if records:
        unknown_count = sum(1 for row in records if _text(row.get("is_in_recovery")) == NOT_COLLECTED)
        if unknown_count:
            return "数据库运行状态检查完成，部分实例角色信息未采集，建议补充核查。"
        return "数据库运行状态检查完成，各实例状态无明显异常。"
    return "数据库运行状态检查结果未采集。"


def _disk_conclusion(data: dict[str, Any]) -> str:
    max_percent = _max_disk_use(data)
    if max_percent is None:
        return "未获取到磁盘空间使用率，建议补充采集。"
    if max_percent < 70:
        return f"当前全集群最高磁盘使用率为 {max_percent:g}%，磁盘空间充足，整体状态为良好。"
    if max_percent < 80:
        return f"当前全集群最高磁盘使用率为 {max_percent:g}%，建议持续关注磁盘增长趋势。"
    return f"当前全集群最高磁盘使用率为 {max_percent:g}%，建议清理历史文件或评估扩容。"


def _ha_conclusion(data: dict[str, Any]) -> str:
    rows = data.get("cluster", {}).get("ha_status", [])
    if not rows:
        return "高可用同步状态检查结果未采集。"
    delayed = [row for row in rows if _number(row.get("pg_xlog_location_diff")) > 0]
    if delayed:
        return "主备同步存在位点差异，建议结合业务写入压力持续关注。"
    return "各备节点与 Master 主节点同步状态正常。"


def _memory_conclusion(data: dict[str, Any]) -> str:
    values = []
    for node in data.get("nodes", []):
        value = _number(node.get("memory", {}).get("usage_percent"))
        if value > 0:
            values.append(value)
    if not values:
        return "内存使用率未采集。"
    max_value = max(values)
    if max_value < 80:
        return f"检查节点最高内存使用率为 {max_value:g}%，内存使用情况良好。"
    if max_value < 90:
        return f"检查节点最高内存使用率为 {max_value:g}%，建议持续观察内存使用趋势。"
    return f"检查节点最高内存使用率为 {max_value:g}%，建议排查内存占用较高的会话或进程。"


def _database_info_conclusion(data: dict[str, Any]) -> str:
    databases = data.get("database", {}).get("databases", [])
    if databases:
        return "数据库信息检查完成，数据库列表、容量和年龄等信息已完成核查。"
    return "数据库信息检查结果未采集。"


def _os_conclusion(data: dict[str, Any]) -> str:
    return "操作系统版本信息检查完成。" if data.get("nodes") else "操作系统版本信息未采集。"


def _database_version_conclusion(data: dict[str, Any]) -> str:
    version = _text(data.get("database", {}).get("version"))
    if version == NOT_COLLECTED:
        return "数据库版本检查结果未采集。"
    return "数据库版本信息检查完成。"


def _cpu_conclusion(data: dict[str, Any]) -> str:
    cores = [node.get("cpu_cores") for node in data.get("nodes", []) if node.get("cpu_cores") not in (None, "", NOT_COLLECTED)]
    return "CPU 核数信息检查完成。" if cores else "CPU 核数信息未采集。"


def _cpu_daily_conclusion(data: dict[str, Any]) -> str:
    cpu_rows = [node.get("cpu_daily", {}) for node in data.get("nodes", [])]
    if not cpu_rows:
        return "CPU 一天使用信息未采集。"
    min_idle = min((_number(row.get("avg_idle")) for row in cpu_rows if _number(row.get("avg_idle")) > 0), default=0)
    max_iowait = max((_number(row.get("max_iowait")) for row in cpu_rows), default=0)
    if min_idle and min_idle < 30:
        return "CPU 空闲率偏低，建议排查高 CPU SQL、后台任务和系统负载。"
    if max_iowait > 10:
        return "IO 等待偏高，建议关注存储响应时间和数据库写入压力。"
    return "CPU 使用率整体平稳，未发现明显 CPU 资源瓶颈。"


def _replication_slot_conclusion(data: dict[str, Any]) -> str:
    slots = data.get("cluster", {}).get("replication_slots", [])
    if not slots:
        return "复制槽状态检查结果未采集。"
    inactive = [slot for slot in slots if str(slot.get("active", "")).lower() in {"false", "0", "f", "no"}]
    delayed = [slot for slot in slots if _number(slot.get("delay_lsn")) > 0]
    if inactive or delayed:
        return "复制槽检查发现需关注项，建议核查复制槽活跃状态和延迟情况。"
    return "复制槽状态正常。"


def _cluster_status_conclusion(data: dict[str, Any]) -> str:
    status = _text(data.get("cluster", {}).get("cluster_status"))
    if status == NOT_COLLECTED:
        return "本次巡检未获取到集群整体状态检查结果，建议后续补充 gs_om -t status 等集群状态采集。"
    if status.lower() == "normal":
        return "集群整体运行状态正常。"
    return "集群整体状态存在需关注项，建议进一步核查。"


def _gs_check_conclusion(data: dict[str, Any]) -> str:
    summary = data.get("cluster", {}).get("gs_check_summary", {})
    ng_count = int(_number(summary.get("ng_count")))
    if ng_count:
        return f"gs_check 巡检存在 {ng_count} 项 NG，建议根据检查结果进行核查和整改。"
    if summary:
        return "gs_check 巡检未发现 NG 项。"
    return "gs_check 巡检信息未采集。"


def _large_table_conclusion(data: dict[str, Any]) -> str:
    count = len(data.get("database", {}).get("large_tables", []))
    if count:
        return f"大表检查发现 {count} 条记录，建议结合业务增长趋势评估归档、分区或历史数据清理策略。"
    return "大表检查未发现需关注记录。"


def _unused_index_conclusion(data: dict[str, Any]) -> str:
    count = len(data.get("database", {}).get("unused_indexes", []))
    if count:
        return f"未使用索引检查发现 {count} 条记录，建议确认后再考虑清理，避免误删业务依赖索引。"
    return "未发现未使用索引。"


def _index_suggestion_conclusion(data: dict[str, Any]) -> str:
    count = len(data.get("database", {}).get("index_suggestions", []))
    if count:
        return f"存在 {count} 条索引优化建议，建议结合高频 SQL、慢 SQL 和业务访问路径评估。"
    return "索引建议检查未发现需关注记录。"


def _table_bloat_conclusion(data: dict[str, Any]) -> str:
    count = len(data.get("database", {}).get("table_bloat", []))
    if count:
        return f"表膨胀检查发现 {count} 条记录，建议结合维护窗口执行 VACUUM / ANALYZE 或表维护操作。"
    return "表膨胀检查未发现需关注记录。"


def _major_findings_sentence(data: dict[str, Any], database: dict[str, Any]) -> str:
    findings: list[str] = []
    if int(_number(data.get("cluster", {}).get("gs_check_summary", {}).get("ng_count"))) > 0:
        findings.append("gs_check 存在 NG 项")
    if database.get("table_bloat"):
        findings.append("部分表存在膨胀")
    if database.get("index_suggestions"):
        findings.append("存在索引优化建议")
    if _text(data.get("cluster", {}).get("cluster_status")) == NOT_COLLECTED:
        findings.append("集群总体状态需补充采集")
    return "；".join(findings) + "。" if findings else "本次巡检未发现明显风险项。"


def _max_disk_use(data: dict[str, Any]) -> float | None:
    values = []
    summary_value = _number(data.get("cluster", {}).get("resource_summary", {}).get("cluster_max_disk_use_percent"))
    if summary_value > 0:
        values.append(summary_value)
    for node in data.get("nodes", []):
        for disk in node.get("disks", []):
            value = _number(disk.get("use_percent"))
            if value > 0:
                values.append(value)
    return max(values) if values else None


def _gs_check_suggestion(check_name: Any) -> str:
    name = str(check_name or "")
    if "CheckDirPermissions" in name:
        return "建议根据安全规范核查目录权限，避免权限过宽。"
    if "CheckGUCValue" in name:
        return "建议核查数据库参数值是否符合当前集群规模和运维规范。"
    if "CheckSysadminUser" in name:
        return "建议核查 sysadmin 用户是否符合最小权限和安全管理要求。"
    if "CheckHashIndex" in name:
        return "建议评估 hash index 使用情况、兼容性和维护风险。"
    return "建议结合 gs_check 检查结果进行核查和整改。"


def _inspection_date(data: dict[str, Any]) -> str:
    value = _text(data.get("report", {}).get("inspection_date"))
    if value != NOT_COLLECTED:
        return value
    return NOT_COLLECTED


def _short_text(value: Any, limit: int) -> str:
    text = " ".join(str(value or "").split())
    if len(text) <= limit:
        return text
    return text[: limit - 1] + "…"


def _text(value: Any, default: str = NOT_COLLECTED) -> str:
    if value is None:
        return default
    text = str(value).strip()
    return text if text else default


def _number(value: Any) -> float:
    try:
        if value is None or value == "":
            return 0.0
        return float(str(value).strip().replace("%", ""))
    except (TypeError, ValueError):
        return 0.0
